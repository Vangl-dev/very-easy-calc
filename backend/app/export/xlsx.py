from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.export.comum import (
    bases_txt,
    critério_encargo_txt,
    criterios_txt,
    dec_memoria,
    pct_memoria,
    regra_parcela_txt,
)

CABECALHO_FILL = PatternFill("solid", fgColor="1F4E79")
CABECALHO_FONT = Font(color="FFFFFF", bold=True)
SECAO_FONT = Font(bold=True)
ROTULO_FILL = PatternFill("solid", fgColor="EEF2F7")
THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ALINHAMENTO = Alignment(vertical="center", wrap_text=True)
ALINHAMENTO_DIR = Alignment(vertical="center", horizontal="right")
MOEDA_HEADERS = {
    "Principal",
    "Salário mínimo",
    "Pagamentos",
    "Diferença",
    "Saldo",
    "Correção",
    "Juros",
    "Multa",
    "Honorários",
    "Total",
    "Total pago",
    "Adicional",
}


def _moeda(valor: object | None) -> float | None:
    d = dec_memoria(valor)
    return None if d is None else float(d)


def _cabecalho(linha: list[str]) -> None:
    for celula in linha:
        celula.fill = CABECALHO_FILL
        celula.font = CABECALHO_FONT
        celula.alignment = ALINHAMENTO
        celula.border = BORDER


def _larguras(ws, larguras: list[float]) -> None:
    for idx, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = largura


def _planilha_resumo(wb: Workbook, memoria: dict) -> None:
    ws = wb.active
    ws.title = "Resumo"
    ident = memoria.get("identificacao") or {}
    resumo = memoria.get("resumo") or {}
    criterios = memoria.get("criterios") or {}
    bases = memoria.get("bases") or []

    ws.append(["Very Easy Calc — memória de cálculo"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(["Identificação"])
    ws["A3"].font = SECAO_FONT
    identificacao = [
        ("Código", memoria.get("codigo") or ""),
        ("Tipo", memoria.get("tipo") or ""),
        ("Título", ident.get("titulo") or ""),
        ("Processo", ident.get("processo") or ""),
        ("Parte", ident.get("parte") or ""),
        ("Aplicativo", ident.get("app_name") or ""),
        ("Versão do app", ident.get("app_version") or ""),
        ("Data de geração", ident.get("data_geracao") or ""),
        ("Data-base", ident.get("data_base") or ""),
        ("Esquema da memória", memoria.get("esquema") or ""),
    ]
    for rotulo, valor in identificacao:
        ws.append([rotulo, valor])
        ws.cell(ws.max_row, 1).fill = ROTULO_FILL
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.cell(ws.max_row, 1).border = BORDER
        ws.cell(ws.max_row, 2).border = BORDER
        ws.cell(ws.max_row, 2).alignment = ALINHAMENTO

    ws.append([])
    ws.append(["Critérios / configurações"])
    ws.cell(ws.max_row, 1).font = SECAO_FONT
    for rotulo, valor in criterios_txt(criterios):
        ws.append([rotulo, valor])
        ws.cell(ws.max_row, 1).fill = ROTULO_FILL
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.cell(ws.max_row, 1).border = BORDER
        ws.cell(ws.max_row, 2).border = BORDER
        ws.cell(ws.max_row, 2).alignment = ALINHAMENTO

    ws.append([])
    ws.append(["Bases e versionamento"])
    ws.cell(ws.max_row, 1).font = SECAO_FONT
    ws.append(["Key", "Versão", "Nome", "Fonte", "Papel", "Cobertura", "URL"])
    _cabecalho(ws[ws.max_row])
    if not bases:
        ws.append(["—", "—", "Nenhuma base indexadora utilizada", "—", "—", "—", "—"])
    for base in bases:
        ws.append(
            [
                base.get("key") or "",
                base.get("versao") or "",
                base.get("nome") or "",
                base.get("fonte") or "",
                base.get("papel") or "",
                base.get("cobertura") or "",
                base.get("fonte_url") or "",
            ]
        )
        for col in range(1, 8):
            ws.cell(ws.max_row, col).border = BORDER
            ws.cell(ws.max_row, col).alignment = ALINHAMENTO

    ws.append([])
    ws.append(["Resumo dos valores"])
    ws.cell(ws.max_row, 1).font = SECAO_FONT
    ws.append(["Campo", "Valor"])
    _cabecalho(ws[ws.max_row])
    campos = [
        ("Parcelas", resumo.get("qtd_parcelas")),
        ("Período início", resumo.get("periodo_inicio")),
        ("Período fim", resumo.get("periodo_fim")),
        ("Principal", _moeda(resumo.get("principal"))),
        ("Correção", _moeda(resumo.get("correcao"))),
        ("Juros", _moeda(resumo.get("juros"))),
        ("Multa", _moeda(resumo.get("multa"))),
        ("Honorários", _moeda(resumo.get("honorarios"))),
        ("Total", _moeda(resumo.get("total"))),
        ("Pagamentos", _moeda(resumo.get("pagamentos"))),
        ("Saldo", _moeda(resumo.get("saldo"))),
        ("Adicional", _moeda(resumo.get("adicional"))),
    ]
    for rotulo, valor in campos:
        ws.append([rotulo, valor])
        ws.cell(ws.max_row, 1).fill = ROTULO_FILL
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.cell(ws.max_row, 1).border = BORDER
        celula = ws.cell(ws.max_row, 2)
        celula.border = BORDER
        if isinstance(valor, (int, float)):
            celula.number_format = 'R$ #,##0.00'
            celula.alignment = ALINHAMENTO_DIR

    avisos = memoria.get("avisos") or []
    if avisos:
        ws.append([])
        ws.append(["Avisos"])
        ws.cell(ws.max_row, 1).font = SECAO_FONT
        for aviso in avisos:
            ws.append([aviso])

    _larguras(ws, [28, 22, 28, 24, 28, 28, 36])


def _planilha_parcelas(wb: Workbook, memoria: dict) -> None:
    ws = wb.create_sheet("Parcelas")
    parcelas = memoria.get("parcelas") or []
    criterios = memoria.get("criterios") or {}
    tem_alimentos = bool(criterios.get("alimentos")) or any(
        p.get("regra") or p.get("tipo_regra") for p in parcelas
    )

    cabecalhos = [
        "Número",
        "Competência",
        "Vencimento",
        "Descrição",
        "Principal",
        "Salário mínimo",
        "Percentual",
        "Regra / detalhe",
        "Pagamentos",
        "Diferença",
        "Saldo",
        "Início correção",
        "Correção",
        "Índice (fator)",
        "Correção aplicada",
        "Início juros",
        "Juros",
        "Juros aplicados",
        "Multa",
        "Multa aplicada",
        "Honorários",
        "Honorários aplicados",
        "Total",
        "Total pago",
        "Adicional",
    ]
    ws.append(cabecalhos)
    _cabecalho(ws[1])

    for parcela in parcelas:
        correcao = parcela.get("correcao") or {}
        juros = parcela.get("juros") or {}
        multa = parcela.get("multa") or {}
        honorarios = parcela.get("honorarios") or {}
        regra = parcela.get("regra") or {}
        if tem_alimentos:
            salario = parcela.get("salario_minimo")
            percentual = parcela.get("percentual")
            detalhe = regra_parcela_txt(parcela) or parcela.get("detalhe") or ""
        else:
            salario = None
            percentual = None
            detalhe = regra_parcela_txt(parcela) or parcela.get("descricao") or ""

        linha = [
            parcela.get("numero"),
            parcela.get("competencia_rotulo") or parcela.get("competencia"),
            parcela.get("vencimento"),
            parcela.get("descricao") or "",
            _moeda(parcela.get("principal")),
            _moeda(salario),
            _moeda(percentual),
            detalhe,
            _moeda(parcela.get("total_pago")),
            _moeda(parcela.get("diferenca")),
            _moeda(parcela.get("saldo")),
            parcela.get("inicio_correcao") or "",
            _moeda(correcao.get("correcao")),
            str(correcao.get("fator_exibicao") or correcao.get("razao") or ""),
            correcao.get("aplicado"),
            parcela.get("inicio_juros") or "",
            _moeda(juros.get("juros")),
            juros.get("aplicado"),
            _moeda(multa.get("valor")),
            multa.get("aplicado"),
            _moeda(honorarios.get("valor")),
            honorarios.get("aplicado"),
            _moeda(parcela.get("total")),
            _moeda(parcela.get("total_pago")),
            _moeda(parcela.get("adicional")),
        ]
        ws.append(linha)
        row = ws.max_row
        for col in range(1, len(cabecalhos) + 1):
            celula = ws.cell(row, col)
            celula.border = BORDER
            celula.alignment = ALINHAMENTO
        for col, header in enumerate(cabecalhos, start=1):
            celula = ws.cell(row, col)
            if header in MOEDA_HEADERS:
                celula.number_format = "R$ #,##0.00"
                celula.alignment = ALINHAMENTO_DIR
            elif header == "Percentual":
                celula.number_format = "0.####"
                celula.alignment = ALINHAMENTO_DIR

    if not parcelas:
        ws.append(["Nenhuma parcela registrada na memória."])

    _larguras(
        ws,
        [10, 14, 14, 34, 14, 16, 12, 26, 14, 14, 14, 14, 14, 14, 14, 14, 12, 14,
         12, 14, 14, 16, 14, 14, 14],
    )
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cabecalhos))}{max(1, ws.max_row)}"
    ws.freeze_panes = "A2"


def gerar_xlsx(memoria: dict) -> bytes:
    """Gera XLSX openpyxl exclusivamente a partir da memória de cálculo."""
    wb = Workbook()
    _planilha_resumo(wb, memoria)
    _planilha_parcelas(wb, memoria)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
