from __future__ import annotations

import io
import json
import re

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from reportlab.platypus import KeepTogether, Paragraph, Table

from app import db
from app.export.comum import brl_memoria, dec_memoria
from app.export.pdf import PAGINA_ALIMENTOS, PAGINA_CIVEL, gerar_pdf, montar_historia
from app.export.xlsx import gerar_xlsx
from app.main import create_app


@pytest.fixture()
def client(tmp_path):
    db.configure(f"sqlite:///{tmp_path / 'test.db'}")
    application = create_app()
    with TestClient(application) as c:
        yield c
    db.reset()


def _payload_alimentos(**kwargs) -> dict:
    base = {
        "tipo": "alimentos",
        "titulo": "Pensão alimentícia",
        "processo": "0001234-56.2021.8.26.0100",
        "parte": "Autor",
        "data_base": "2026-08-31",
        "alimentos": {
            "regras": [
                {
                    "inicio": "01/2021",
                    "fim": "12/2021",
                    "tipo": "salario_minimo",
                    "percentual": "0.30",
                }
            ],
            "ate": "12/2021",
            "dia_vencimento": 10,
        },
        "correcao": {"aplicar": False},
        "juros": {"aplicar": False},
        "multa": {"aplicar": False},
        "honorarios": {"aplicar": False},
        "pagamentos": [
            {
                "competencia": "2021-01",
                "data": "2026-01-10",
                "valor": "100.00",
                "descricao": "depósito",
            }
        ],
    }
    base.update(kwargs)
    return base


def _memoria_fixa() -> dict:
    return {
        "esquema": "memoria-calculo/1.0",
        "codigo": "VEC-2026-000001",
        "tipo": "alimentos",
        "identificacao": {
            "titulo": "Pensão alimentícia",
            "processo": "0001234-56.2021.8.26.0100",
            "parte": "Autor",
            "app_name": "Very Easy Calc",
            "app_version": "1.0.0",
            "data_geracao": "2026-10-01T12:00:00+00:00",
            "data_base": "2026-08-31",
        },
        "criterios": {
            "data_base": "2026-08-31",
            "correcao": {"aplicar": False, "indexador": "tjsp_tabela_pratica"},
            "juros": {"aplicar": False, "modalidade": "taxa_legal"},
            "multa": {"aplicar": False},
            "honorarios": {"aplicar": False},
            "alimentos": {
                "regras": [
                    {
                        "inicio": "01/2021",
                        "fim": "12/2021",
                        "tipo": "salario_minimo",
                        "percentual": "0.30",
                    }
                ],
                "ate": "12/2021",
                "dia_vencimento": 10,
            },
            "pagamentos_informados": [
                {"competencia": "2021-01", "data": "2026-01-10", "valor": "100.00"}
            ],
        },
        "bases": [
            {
                "key": "salario_minimo",
                "nome": "Salário mínimo",
                "versao": "2026.12",
                "fonte": "Ipeadata / Ministério da Economia",
                "fonte_url": "https://ipeadata.gov.br",
                "papel": "Composição do principal",
                "cobertura": "07/1994 a 12/2026",
            }
        ],
        "resumo": {
            "qtd_parcelas": 2,
            "periodo_inicio": "2021-01-10",
            "periodo_fim": "2021-02-10",
            "principal": "660.00",
            "correcao": "0.00",
            "juros": "0.00",
            "multa": "0.00",
            "honorarios": "0.00",
            "total": "660.00",
            "pagamentos": "100.00",
            "saldo": "560.00",
            "adicional": "0.00",
        },
        "parcelas": [
            {
                "numero": 1,
                "competencia": "2021-01",
                "competencia_rotulo": "01/2021",
                "vencimento": "2021-01-10",
                "descricao": "30% do salário mínimo (30% × R$ 1.100,00)",
                "principal": "330.00",
                "salario_minimo": "1100.00",
                "percentual": "0.30",
                "tipo_regra": "salario_minimo",
                "origem": "30% do salário mínimo",
                "detalhe": "30% × R$ 1.100,00",
                "regra": {
                    "tipo": "salario_minimo",
                    "origem": "30% do salário mínimo",
                    "detalhe": "30% × R$ 1.100,00",
                    "salario_minimo": "1100.00",
                    "percentual": "0.30",
                },
                "correcao": {
                    "aplicado": False,
                    "principal": "330.00",
                    "valor_corrigido": "330.00",
                    "correcao": "0.00",
                    "motivo": "Correção monetária não aplicada",
                },
                "juros": {
                    "aplicado": False,
                    "base": "330.00",
                    "incidencia": "valor corrigido",
                    "modalidade": "taxa_legal",
                    "juros": "0.00",
                    "motivo": "Juros não aplicados",
                },
                "multa": {"aplicado": False, "valor": "0.00", "motivo": "Multa não aplicado"},
                "honorarios": {
                    "aplicado": False,
                    "valor": "0.00",
                    "motivo": "Honorários não aplicado",
                },
                "pagamentos": [
                    {"data": "2026-01-10", "valor": "100.00", "descricao": "depósito"}
                ],
                "total": "330.00",
                "total_pago": "100.00",
                "saldo": "230.00",
                "adicional": "0.00",
            },
            {
                "numero": 2,
                "competencia": "2021-02",
                "competencia_rotulo": "02/2021",
                "vencimento": "2021-02-10",
                "descricao": "30% do salário mínimo (30% × R$ 1.100,00)",
                "principal": "330.00",
                "salario_minimo": "1100.00",
                "percentual": "0.30",
                "tipo_regra": "salario_minimo",
                "origem": "30% do salário mínimo",
                "detalhe": "30% × R$ 1.100,00",
                "regra": {
                    "tipo": "salario_minimo",
                    "origem": "30% do salário mínimo",
                    "detalhe": "30% × R$ 1.100,00",
                    "salario_minimo": "1100.00",
                    "percentual": "0.30",
                },
                "correcao": {
                    "aplicado": False,
                    "principal": "330.00",
                    "valor_corrigido": "330.00",
                    "correcao": "0.00",
                    "motivo": "Correção monetária não aplicada",
                },
                "juros": {
                    "aplicado": False,
                    "base": "330.00",
                    "incidencia": "valor corrigido",
                    "modalidade": "taxa_legal",
                    "juros": "0.00",
                    "motivo": "Juros não aplicados",
                },
                "multa": {"aplicado": False, "valor": "0.00", "motivo": "Multa não aplicado"},
                "honorarios": {
                    "aplicado": False,
                    "valor": "0.00",
                    "motivo": "Honorários não aplicado",
                },
                "pagamentos": [],
                "total": "330.00",
                "total_pago": "0.00",
                "saldo": "330.00",
                "adicional": "0.00",
            },
        ],
        "avisos": [],
    }


def test_export_pdf_inexistente_404(client: TestClient) -> None:
    assert client.get("/api/calculos/VEC-2026-000099/pdf").status_code == 404
    assert client.get("/api/calculos/VEC-2026-000099/xlsx").status_code == 404


def test_pdf_gerado_e_nao_vazio(client: TestClient) -> None:
    codigo = client.post("/api/calculos", json=_payload_alimentos()).json()["codigo"]
    response = client.get(f"/api/calculos/{codigo}/pdf")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"].endswith(f'filename="{codigo}.pdf"')
    conteudo = response.content
    assert len(conteudo) > 500
    assert conteudo.startswith(b"%PDF")


def test_pdf_contem_cabecalho_e_memoria(client: TestClient) -> None:
    memoria = _memoria_fixa()
    pdf = gerar_pdf(memoria)
    assert pdf.startswith(b"%PDF")
    assert len(pdf) > 800
    # ReportLab grava metadados legíveis no PDF
    assert b"Very Easy Calc" in pdf or b"VEC-2026-000001" in pdf


def test_pdf_escolhe_orientacao_automaticamente_por_modulo() -> None:
    alimentos = _memoria_fixa()
    civel = json.loads(json.dumps(alimentos))
    civel["tipo"] = "parcelas"
    civel["criterios"]["alimentos"] = None

    assert PAGINA_CIVEL[0] < PAGINA_CIVEL[1]
    assert PAGINA_ALIMENTOS[0] > PAGINA_ALIMENTOS[1]
    assert gerar_pdf(civel).startswith(b"%PDF")
    assert gerar_pdf(alimentos).startswith(b"%PDF")

    texto_civel = _texto_historia(montar_historia(civel))
    assert "Parcelas do cálculo — valores principais" in texto_civel
    assert "Parcelas do cálculo — composição e saldo" in texto_civel


def test_metodologia_numera_secoes_sem_lacuna_quando_taxa_legal_nao_se_aplica() -> None:
    texto = _texto_historia(montar_historia(_memoria_fixa()))

    assert "1. Correção monetária" in texto
    assert "2. Juros" in texto
    assert "3. Multas e honorários" in texto
    assert "4. Pagamentos e abatimentos" in texto
    assert "5. Fontes, rastreabilidade e arredondamento" in texto
    assert "Taxa Legal e regimes temporais" not in texto


def test_pdf_muitas_parcelas_gera_documento(client: TestClient) -> None:
    payload = _payload_alimentos(
        alimentos={
            "regras": [
                {
                    "inicio": "01/2021",
                    "fim": None,
                    "tipo": "salario_minimo",
                    "percentual": "0.30",
                }
            ],
            "ate": "09/2026",
        },
        pagamentos=[],
    )
    corpo = client.post("/api/calculos", json=payload).json()
    assert corpo["resumo"]["qtd_parcelas"] == 69
    pdf = client.get(f"/api/calculos/{corpo['codigo']}/pdf").content
    assert pdf.startswith(b"%PDF")
    assert len(pdf) > 5000
    memoria_pdf = gerar_pdf(corpo["memoria"])
    assert memoria_pdf.startswith(b"%PDF")
    assert len(memoria_pdf) > 5000


def test_xlsx_gerado_e_nao_vazio(client: TestClient) -> None:
    codigo = client.post("/api/calculos", json=_payload_alimentos()).json()["codigo"]
    response = client.get(f"/api/calculos/{codigo}/xlsx")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert response.headers["content-disposition"].endswith(f'filename="{codigo}.xlsx"')
    conteudo = response.content
    assert len(conteudo) > 800
    assert conteudo[:2] == b"PK"


def test_xlsx_planilhas_e_valores_da_memoria() -> None:
    memoria = _memoria_fixa()
    xlsx = gerar_xlsx(memoria)
    wb = load_workbook(io.BytesIO(xlsx))
    assert wb.sheetnames == ["Resumo", "Parcelas"]

    resumo = wb["Resumo"]
    valores_resumo = {}
    for linha in resumo.iter_rows(min_row=1, max_col=2, values_only=True):
        if linha[0] and linha[1] is not None:
            valores_resumo[str(linha[0])] = linha[1]
    assert valores_resumo["Código"] == "VEC-2026-000001"
    assert valores_resumo["Versão do app"] == "1.0.0"
    assert valores_resumo["Data-base"] == "2026-08-31"
    assert valores_resumo["Principal"] == 660.0
    assert valores_resumo["Total"] == 660.0
    assert valores_resumo["Saldo"] == 560.0
    assert valores_resumo["Pagamentos"] == 100.0
    assert valores_resumo["Parcelas"] == 2

    parcelas = wb["Parcelas"]
    cabecalhos = [c.value for c in next(parcelas.iter_rows(min_row=1, max_row=1))]
    assert "Salário mínimo" in cabecalhos
    assert "Percentual" in cabecalhos
    assert "Principal" in cabecalhos
    assert "Saldo" in cabecalhos
    assert "Total" in cabecalhos

    primeira = [c.value for c in next(parcelas.iter_rows(min_row=2, max_row=2))]
    mapa = dict(zip(cabecalhos, primeira, strict=False))
    assert mapa["Competência"] == "01/2021"
    assert float(mapa["Salário mínimo"]) == 1100.0
    assert float(mapa["Percentual"]) == 0.3
    assert float(mapa["Principal"]) == 330.0
    assert float(mapa["Pagamentos"]) == 100.0
    assert float(mapa["Saldo"]) == 230.0
    assert float(mapa["Total"]) == 330.0
    assert isinstance(mapa["Principal"], (int, float))
    assert isinstance(mapa["Salário mínimo"], (int, float))
    assert not isinstance(mapa["Principal"], str)
    assert not isinstance(mapa["Salário mínimo"], str)
    # percentual fracionado exibido como porcentagem
    assert "30%" in str(mapa["Regra / detalhe"])
    assert "0.30%" not in str(mapa["Regra / detalhe"])
    # memória antiga (1.0) não tinha início de correção/juros nem diferença
    assert not mapa["Início correção"]
    assert not mapa["Início juros"]
    assert mapa["Diferença"] is None


def test_xlsx_api_valores_correspondem_a_memoria(client: TestClient) -> None:
    corpo = client.post("/api/calculos", json=_payload_alimentos()).json()
    memoria = corpo["memoria"]
    xlsx = client.get(f"/api/calculos/{corpo['codigo']}/xlsx").content
    wb = load_workbook(io.BytesIO(xlsx))
    parcelas = wb["Parcelas"]
    cabecalhos = [c.value for c in next(parcelas.iter_rows(min_row=1, max_row=1))]
    linhas = list(parcelas.iter_rows(min_row=2, values_only=True))
    assert len(linhas) == memoria["resumo"]["qtd_parcelas"]
    for memoria_p, linha in zip(memoria["parcelas"], linhas, strict=True):
        mapa = dict(zip(cabecalhos, linha, strict=False))
        assert mapa["Competência"] == memoria_p["competencia_rotulo"]
        assert mapa["Principal"] == float(memoria_p["principal"])
        assert mapa["Salário mínimo"] == float(memoria_p["salario_minimo"])
        assert mapa["Percentual"] == float(memoria_p["percentual"])
        assert mapa["Saldo"] == float(memoria_p["saldo"])
        assert mapa["Total"] == float(memoria_p["total"])
        assert mapa["Pagamentos"] == float(memoria_p["total_pago"])
        assert mapa["Diferença"] == float(memoria_p["diferenca"])
        assert (mapa["Início correção"] or "") == (memoria_p["inicio_correcao"] or "")
        assert (mapa["Início juros"] or "") == (memoria_p["inicio_juros"] or "")


def test_xlsx_muitas_parcelas(client: TestClient) -> None:
    payload = _payload_alimentos(
        alimentos={
            "regras": [
                {
                    "inicio": "01/2021",
                    "fim": None,
                    "tipo": "salario_minimo",
                    "percentual": "0.30",
                }
            ],
            "ate": "09/2026",
        },
        pagamentos=[],
    )
    corpo = client.post("/api/calculos", json=payload).json()
    xlsx = client.get(f"/api/calculos/{corpo['codigo']}/xlsx").content
    wb = load_workbook(io.BytesIO(xlsx))
    parcelas = wb["Parcelas"]
    linhas = [linha for linha in parcelas.iter_rows(min_row=2, values_only=True) if linha[0]]
    assert len(linhas) == 69
    assert linhas[0][0] == 1
    assert linhas[-1][0] == 69


def test_export_nao_recalcula_valores(client: TestClient) -> None:
    """Exportação usa a memória; valores não divergem do cálculo."""
    corpo = client.post("/api/calculos", json=_payload_alimentos()).json()
    memoria = corpo["memoria"]
    assert memoria["parcelas"][0]["principal"] == "330.00"
    assert memoria["parcelas"][0]["total"] == "330.00"
    assert memoria["parcelas"][0]["saldo"] == "230.00"

    xlsx = gerar_xlsx(memoria)
    wb = load_workbook(io.BytesIO(xlsx))
    parcelas = wb["Parcelas"]
    cabecalhos = [c.value for c in next(parcelas.iter_rows(min_row=1, max_row=1))]
    primeira = dict(
        zip(cabecalhos, next(parcelas.iter_rows(min_row=2, max_row=2, values_only=True)), strict=False)
    )
    assert primeira["Principal"] == float(memoria["parcelas"][0]["principal"])
    assert primeira["Saldo"] == float(memoria["parcelas"][0]["saldo"])
    assert brl_memoria(memoria["parcelas"][0]["principal"]) == "R$ 330,00"
    assert dec_memoria(memoria["parcelas"][0]["principal"]) is not None


def test_comum_converte_sem_recalcular() -> None:
    assert dec_memoria("330.00") is not None
    assert str(dec_memoria("330.00")) == "330.00"
    assert brl_memoria("1100.00") == "R$ 1.100,00"
    assert brl_memoria(None) == ""


def _payload_civel() -> dict:
    def verba(
        id_: str,
        nome: str,
        *,
        tipo: str = "valor",
        valor: str | None = None,
        base: str | None = None,
        pct: str | None = None,
    ) -> dict:
        return {
            "id": id_,
            "tipo": tipo,
            "nome": nome,
            "valor": valor,
            "data_base": "2022-08-30",
            "descricao": "",
            "base_verba_id": base,
            "percentual": pct,
            "correcao": {
                "aplicar": True,
                "indexador": "tjsp_tabela_pratica",
                "termo_inicial": "2022-08-30",
                "termo_final": None,
            },
            "juros": {
                "aplicar": True,
                "modalidade": "taxa_legal",
                "taxa": None,
                "unidade": "mensal",
                "base": "corrigido",
                "inicio": "2026-04-24",
                "fim": None,
            },
            "multa": {"aplicar": False},
            "honorarios": {"aplicar": False},
            "pagamentos": [],
            "parcelas": [],
        }

    return {
        "tipo": "parcelas",
        "titulo": "Ação de indenização",
        "processo": "0000000-00.0000.0.00.0000",
        "parte": "Autor",
        "data_base": "2026-09-30",
        "verbas": [
            verba("causa", "Valor da causa", valor="44000,00"),
            verba(
                "sucumb",
                "Honorários de sucumbência",
                tipo="percentual",
                valor=None,
                base="causa",
                pct="0.12",
            ),
        ],
        "correcao": {"aplicar": False},
        "juros": {"aplicar": False},
        "multa": {"aplicar": False},
        "honorarios": {"aplicar": False},
        "pagamentos": [],
    }


def _texto_historia(fluxo) -> str:
    partes: list[str] = []
    for item in fluxo:
        if isinstance(item, KeepTogether):
            partes.append(_texto_historia(item._content))
        elif isinstance(item, Paragraph):
            partes.append(item.getPlainText())
        elif isinstance(item, Table):
            for linha in item._cellvalues:
                for celula in linha:
                    if isinstance(celula, Paragraph):
                        partes.append(celula.getPlainText())
                    elif isinstance(celula, (list, tuple)):
                        partes.extend(
                            p.getPlainText()
                            for p in celula
                            if isinstance(p, Paragraph)
                        )
    return "\n".join(partes)


def test_pdf_civel_secoes_e_datas_em_formato_brasileiro(client: TestClient) -> None:
    corpo = client.post("/api/calculos", json=_payload_civel()).json()
    memoria = corpo["memoria"]
    texto = _texto_historia(montar_historia(memoria))

    assert "VERY EASY CALC" in texto
    assert "Memória de Cálculo" in texto
    assert memoria["codigo"] in texto
    # datas entregues em formato brasileiro, sem ISO 8601
    assert "30/09/2026" in texto
    assert "30/08/2022" in texto
    assert "24/04/2026" in texto
    assert re.search(r"\b\d{4}-\d{2}-\d{2}\b", texto) is None

    # hierarquia: resumo -> configuração -> memória -> critérios -> bases
    ordem = [
        "Identificação",
        "Resumo do cálculo",
        "Configuração das verbas",
        "Memória de cálculo",
        "Critérios",
        "Bases e versionamento",
    ]
    posicoes = [texto.index(rotulo) for rotulo in ordem]
    assert posicoes == sorted(posicoes)
    assert texto.index("Resumo do cálculo") < texto.index("Memória de cálculo")
    # detalhes técnicos ficam no fim, não no topo
    assert "Esquema da memória" not in texto
    assert "base_verba_id" not in texto


def test_pdf_resumo_por_verba_usa_valores_ja_calculados(client: TestClient) -> None:
    memoria = client.post("/api/calculos", json=_payload_civel()).json()["memoria"]
    texto = _texto_historia(montar_historia(memoria))

    causa, sucumb = memoria["parcelas"]
    assert brl_memoria(causa["principal"]) == "R$ 44.000,00"
    assert brl_memoria(sucumb["principal"]) == "R$ 5.280,00"

    for parcela in memoria["parcelas"]:
        correcao = parcela["correcao"]["correcao"]
        juros = parcela["juros"]["juros"]
        assert brl_memoria(correcao) in texto
        assert brl_memoria(juros) in texto

    assert "VALOR DA CAUSA" in texto
    assert "HONORÁRIOS DE SUCUMBÊNCIA" in texto
    assert "12% sobre Valor da causa" in texto
    assert "TOTAL GERAL" in texto
    assert brl_memoria(memoria["resumo"]["total"]) in texto


def test_pdf_deixa_taxa_legal_explicita_quando_efetivamente_usada(
    client: TestClient,
) -> None:
    memoria = client.post("/api/calculos", json=_payload_civel()).json()["memoria"]
    texto = _texto_historia(montar_historia(memoria))

    juros = memoria["parcelas"][0]["juros"]
    assert juros["aplicado"] is True
    assert juros["indexador"] == "taxa_legal"

    assert "Série aplicada" in texto
    assert "Taxa Legal (BCB SGS 29543)" in texto
    assert "Banco Central do Brasil" in texto
    assert "Taxa Legal variável" in texto
    # período efetivamente percorrido pelos trechos do motor
    assert "24/04/2026 a 30/09/2026" in texto
    assert str(juros["trechos"][0]["dias"]) in texto


def test_pdf_nao_afirma_taxa_legal_quando_nao_foi_aplicada() -> None:
    memoria = _memoria_fixa()
    assert memoria["parcelas"][0]["juros"]["aplicado"] is False
    texto = _texto_historia(montar_historia(memoria))

    assert "não aplicados" in texto
    assert "Série aplicada" not in texto
    assert "Taxa Legal (BCB SGS 29543)" not in texto


def test_pdf_metodologia_taxa_legal_e_rastreabilidade_derivadas_da_memoria(
    client: TestClient,
) -> None:
    memoria = client.post("/api/calculos", json=_payload_civel()).json()["memoria"]
    texto = _texto_historia(montar_historia(memoria))

    inicio = texto.index("Metodologia de Cálculo e Fundamentação Normativa")
    assert inicio > texto.index("Bases e versionamento")
    secao = texto[inicio:]
    assert "Lei nº 14.905/2024" in secao
    assert "Resolução CMN nº 5.171/2024" in secao
    assert "resultado negativo é tratado como zero" in secao
    assert "juros simples" in secao
    assert "Banco Central do Brasil" in secao
    assert "versão" in secao


def test_pdf_metodologia_nao_inventa_encargos_em_memoria_legada() -> None:
    memoria = _memoria_fixa()
    texto = _texto_historia(montar_historia(memoria))
    inicio = texto.index("Metodologia de Cálculo e Fundamentação Normativa")
    secao = texto[inicio:]

    assert "Não houve correção monetária efetivamente aplicada" in secao
    assert "Não houve juros efetivamente aplicados" in secao
    assert "Multa\nnão aplicado nas parcelas registradas" in secao
    assert "Honorários\nnão aplicado nas parcelas registrados" not in secao
    assert "Honorários\nnão aplicado nas parcelas registradas" in secao
    assert "Lei nº 14.905/2024" not in secao


def test_pdf_metodologia_descreve_criterio_fixo_e_abatimento_sem_taxa_legal(
    client: TestClient,
) -> None:
    memoria = client.post(
        "/api/calculos",
        json=_payload_alimentos(
            correcao={"aplicar": False},
            juros={"aplicar": True, "modalidade": "fixa", "taxa": "0.01", "unidade": "mensal"},
            multa={"aplicar": True, "percentual": "0.02", "base": "corrigido"},
            honorarios={"aplicar": True, "percentual": "0.10", "base": "principal"},
        ),
    ).json()["memoria"]
    texto = _texto_historia(montar_historia(memoria))
    secao = texto[texto.index("Metodologia de Cálculo e Fundamentação Normativa"):]

    assert "Taxa fixa" in secao
    assert "por verba/parcela" in secao
    assert "Cada pagamento foi somado" in secao
    assert "Lei nº 14.905/2024" not in secao


def _payload_alimentos_divida() -> dict:
    """Dívida 11/2025→03/2026 atualizada até 30/09/2026 (padrão do formulário)."""
    return _payload_alimentos(
        titulo="Pensão alimentícia — dívida",
        data_base="2026-09-30",
        alimentos={
            "regras": [
                {
                    "inicio": "11/2025",
                    "fim": "03/2026",
                    "tipo": "salario_minimo",
                    "percentual": "0.35",
                }
            ],
            "ate": "03/2026",
            "dia_vencimento": 10,
        },
        pagamentos=[],
        correcao={"aplicar": True, "indexador": "tjsp_tabela_pratica"},
        juros={
            "aplicar": True,
            "modalidade": "fixa",
            "taxa": "0.01",
            "unidade": "mensal",
            "base": "corrigido",
        },
    )


def _tabelas(fluxo) -> list[Table]:
    tabelas: list[Table] = []
    for item in fluxo:
        if isinstance(item, KeepTogether):
            tabelas.extend(_tabelas(item._content))
        elif isinstance(item, Table):
            tabelas.append(item)
    return tabelas


def _titulo_tabela(tabela: Table) -> str:
    primeira = tabela._cellvalues[0][0]
    return primeira.getPlainText() if isinstance(primeira, Paragraph) else ""


def _cabecalhos_tabela(tabela: Table) -> list[str]:
    return [
        celula.getPlainText()
        for celula in tabela._cellvalues[1]
        if isinstance(celula, Paragraph)
    ]


def test_alimentos_periodo_da_divida_nao_e_a_data_de_atualizacao(
    client: TestClient,
) -> None:
    corpo = client.post("/api/calculos", json=_payload_alimentos_divida()).json()
    memoria = corpo["memoria"]
    parcelas = memoria["parcelas"]

    # cinco competências, todas atualizadas até 30/09/2026
    assert [p["competencia_rotulo"] for p in parcelas] == [
        "11/2025",
        "12/2025",
        "01/2026",
        "02/2026",
        "03/2026",
    ]
    assert {p["correcao"]["termo_final"] for p in parcelas} == {"2026-09"}
    assert {
        trecho["fim"] for p in parcelas for trecho in p["juros"]["trechos"]
    } == {"2026-09-30"}

    texto = _texto_historia(montar_historia(memoria))
    # março/2026 é o fim da dívida; setembro/2026 é a data de atualização
    assert "11/2025 a 03/2026" in texto
    assert "Atualizado até" in texto
    assert "30/09/2026" in texto


def test_pdf_alimentos_tem_estrutura_de_leitura_humana(client: TestClient) -> None:
    corpo = client.post("/api/calculos", json=_payload_alimentos_divida()).json()
    memoria = corpo["memoria"]
    texto = _texto_historia(montar_historia(memoria))

    assert "MEMÓRIA DE CÁLCULO — ALIMENTOS" in texto
    for secao in (
        "Resultado",
        "Regra alimentar",
        "Atualização",
        "Resumo final",
        "Bases e versionamento",
    ):
        assert secao in texto
    assert texto.index("Resultado") < texto.index("Bases e versionamento")

    # datas brasileiras e valores da memória, sem recálculo
    assert re.search(r"\b\d{4}-\d{2}-\d{2}\b", texto) is None
    assert "10/11/2025" in texto
    assert brl_memoria(memoria["resumo"]["principal"]) in texto
    assert brl_memoria(memoria["resumo"]["saldo"]) in texto
    assert brl_memoria(memoria["parcelas"][0]["principal"]) in texto
    assert "TOTAL DO DÉBITO ATUALIZADO" in texto


def test_pdf_alimentos_tabelas_estreitas_sem_coluna_tecnica(client: TestClient) -> None:
    corpo = client.post("/api/calculos", json=_payload_alimentos_divida()).json()
    fluxo = montar_historia(corpo["memoria"])
    tabelas = _tabelas(fluxo)

    principal = next(t for t in tabelas if _titulo_tabela(t) == "Dívida por competência")
    assert _cabecalhos_tabela(principal) == [
        "Competência",
        "Vencimento",
        "Salário mín.",
        "%",
        "Devido",
        "Pago",
        "Saldo",
    ]

    detalhamento = next(
        t for t in tabelas if _titulo_tabela(t) == "Detalhamento da atualização"
    )
    assert _cabecalhos_tabela(detalhamento) == [
        "Competência",
        "Correção",
        "Índice",
        "Juros",
        "Total atualizado",
    ]

    # nenhuma tabela de Alimentos passa de 7 colunas
    com_cabecalho = [t for t in tabelas if len(t._cellvalues) > 1]
    assert max(len(_cabecalhos_tabela(t)) for t in com_cabecalho) <= 7


# ---------------------------------------------------------------------------
# Resumo para petição (Cível e Alimentos)
# ---------------------------------------------------------------------------


def _secao_peticao(texto: str, *, fim: str = "Configuração das verbas") -> str:
    inicio = texto.index("Resumo para petição")
    return texto[inicio : texto.index(fim, inicio)]


def _payload_civel_com_pagamento() -> dict:
    payload = _payload_civel()
    payload["titulo"] = "Execução com pagamento parcial"
    # pagamento na verba que compõe o total (a verba-base é referência)
    payload["verbas"][1]["pagamentos"] = [
        {"data": "2025-06-10", "valor": "1500.00", "descricao": "entrada"}
    ]
    return payload


def test_pdf_civel_resumo_para_peticao_usa_totais_da_memoria(client: TestClient) -> None:
    """Total da execução é o da memória; verba-base fica fora do total."""
    memoria = client.post("/api/calculos", json=_payload_civel()).json()["memoria"]
    texto = _texto_historia(montar_historia(memoria))
    secao = _secao_peticao(texto)
    resumo = memoria["resumo"]

    assert "Resumo para petição" in texto
    assert "TOTAL DA EXECUÇÃO" not in secao
    assert "Saldo final" in secao
    assert "Multa" in secao
    assert "Honorários" in secao
    assert "Pagamentos / abatimentos" in secao
    # mesma semântica do motor: a verba-base não compõe o total
    assert float(resumo["total"]) == pytest.approx(
        sum(float(p["total"]) for p in memoria["parcelas"][1:])
    )
    assert brl_memoria(resumo["principal"]) in secao
    assert brl_memoria(resumo["correcao"]) in secao
    assert brl_memoria(resumo["juros"]) in secao
    assert brl_memoria(resumo["total"]) in secao
    assert "não compõe o total da execução" in texto

    # bloco textual selecionável: data de atualização e total apurado
    assert f"Total apurado: {brl_memoria(resumo['total'])}." in texto
    assert "atualizado até 30/09/2026" in texto
    assert "valores principais" in texto
    assert "composição e saldo" in texto


def test_pdf_civel_peticao_mostra_pagamentos_e_saldo_quando_houver(
    client: TestClient,
) -> None:
    memoria = client.post("/api/calculos", json=_payload_civel_com_pagamento()).json()[
        "memoria"
    ]
    texto = _texto_historia(montar_historia(memoria))
    secao = _secao_peticao(texto)
    resumo = memoria["resumo"]

    assert resumo["pagamentos"] == "1500.00"
    assert brl_memoria(resumo["pagamentos"]) in secao
    assert brl_memoria(resumo["saldo"]) in secao
    assert (
        f"Pagamentos considerados: {brl_memoria(resumo['pagamentos'])}; "
        f"saldo devedor: {brl_memoria(resumo['saldo'])}."
    ) in texto
    # pagamento não é subtraído de novo: saldo vem da memória
    assert float(resumo["saldo"]) == pytest.approx(
        float(resumo["total"]) - float(resumo["pagamentos"])
    )


def test_pdf_civel_peticao_omite_pagamentos_quando_nao_existem(
    client: TestClient,
) -> None:
    memoria = client.post("/api/calculos", json=_payload_civel()).json()["memoria"]
    texto = _texto_historia(montar_historia(memoria))
    assert "Pagamentos considerados" not in _secao_peticao(texto)
    assert "Pagamentos considerados" not in texto


def test_pdf_civel_peticao_com_celulas_rotulo_valores(client: TestClient) -> None:
    """Sem verbas, o resumo para petição vira pares rotulo/valor."""
    payload = _payload_civel()
    payload["verbas"] = []
    payload["parcelas"] = [
        {
            "numero": 1,
            "vencimento": "2024-01-15",
            "principal": "1000.00",
            "descricao": "parcela única",
            "correcao": {"aplicar": False},
            "juros": {"aplicar": False},
            "multa": {"aplicar": False},
            "honorarios": {"aplicar": False},
            "pagamentos": [],
        }
    ]
    memoria = client.post("/api/calculos", json=payload).json()["memoria"]
    texto = _texto_historia(montar_historia(memoria))
    secao = _secao_peticao(texto, fim="Memória de cálculo")

    assert "TOTAL DA EXECUÇÃO" not in secao
    assert "Total atualizado" in secao
    assert brl_memoria(memoria["resumo"]["total"]) in secao
    assert f"Total apurado: {brl_memoria(memoria['resumo']['total'])}." in texto


def test_pdf_alimentos_resumo_para_peticao_compacto(client: TestClient) -> None:
    """Bloco de petição com período, valores da memória e saldo — sem competências."""
    payload = _payload_alimentos_divida()
    payload["pagamentos"] = [
        {"competencia": "2026-01", "data": "2026-01-10", "valor": "200.00"}
    ]
    memoria = client.post("/api/calculos", json=payload).json()["memoria"]
    texto = _texto_historia(montar_historia(memoria))
    inicio = texto.index("Resumo para petição")
    secao = texto[inicio : texto.index("Bases e versionamento", inicio)]
    resumo = memoria["resumo"]

    assert "Período do débito" in secao
    assert "11/2025 a 03/2026" in secao
    assert brl_memoria(resumo["principal"]) in secao
    assert brl_memoria(resumo["correcao"]) in secao
    assert brl_memoria(resumo["juros"]) in secao
    assert brl_memoria(resumo["pagamentos"]) in secao
    assert brl_memoria(resumo["total"]) in secao
    assert "Saldo devedor atualizado" in secao
    assert brl_memoria(resumo["saldo"]) in secao
    assert "30/09/2026" in secao
    # o resumo não repete as competências da tabela detalhada
    assert "02/2026" not in secao
    assert "Salário mín." not in secao

    # bloco textual selecionável, no formato de petição
    assert "Resumo do débito alimentar: período de 11/2025 a 03/2026." in texto
    assert (
        f"Saldo devedor atualizado até 30/09/2026: {brl_memoria(resumo['saldo'])}."
    ) in texto
    # convenção da memória: total e saldo já consideram os pagamentos
    assert float(resumo["saldo"]) == pytest.approx(
        float(resumo["total"]) - float(resumo["pagamentos"])
    )


def test_pdf_alimentos_peticao_omite_campos_nao_aplicados(client: TestClient) -> None:
    payload = _payload_alimentos(
        titulo="Pensão sem atualização",
        correcao={"aplicar": False},
        juros={"aplicar": False},
        pagamentos=[],
    )
    memoria = client.post("/api/calculos", json=payload).json()["memoria"]
    texto = _texto_historia(montar_historia(memoria))
    inicio = texto.index("Resumo para petição")
    secao = texto[inicio : texto.index("Bases e versionamento", inicio)]

    assert "Correção monetária" not in secao
    assert "Juros" not in secao
    assert "Pagamentos / abatimentos" not in secao
    assert "Período do débito" in secao
    assert "Saldo devedor atualizado" in secao
    assert brl_memoria(memoria["resumo"]["saldo"]) in secao


def test_gerar_pdf_nao_altera_a_memoria(client: TestClient) -> None:
    """A exportação lê a memória; nenhum valor matemático pode mudar."""
    memoria = client.post("/api/calculos", json=_payload_civel()).json()["memoria"]
    antes = json.dumps(memoria, sort_keys=True, ensure_ascii=False)
    gerar_pdf(memoria)
    depois = json.dumps(memoria, sort_keys=True, ensure_ascii=False)
    assert antes == depois
