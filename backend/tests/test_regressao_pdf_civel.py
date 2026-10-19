from __future__ import annotations

from datetime import datetime, timezone
from copy import deepcopy

from io import BytesIO

from openpyxl import load_workbook
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, Paragraph, Table

from app.engine.memoria import montar_memoria
from app.export.comum import criterios_txt
from app.export.pdf import LARGURA, montar_historia
from app.export.xlsx import gerar_xlsx


def _texto(fluxo) -> str:
    partes: list[str] = []
    for item in fluxo:
        if isinstance(item, KeepTogether):
            partes.append(_texto(item._content))
        elif isinstance(item, Paragraph):
            partes.append(item.getPlainText())
        elif isinstance(item, Table):
            for linha in item._cellvalues:
                for celula in linha:
                    if isinstance(celula, Paragraph):
                        partes.append(celula.getPlainText())
                    elif isinstance(celula, (list, tuple)):
                        partes.extend(
                            p.getPlainText() for p in celula if isinstance(p, Paragraph)
                        )
        elif isinstance(item, (list, tuple)):
            partes.append(_texto(item))
    return "\n".join(partes)


def _memoria_civel_com_encargos_e_abatimento() -> dict:
    verbas = [
        {
            "id": "entrada",
            "nome": "Entrada",
            "tipo": "valor",
            "valor": "36800.00",
            "data_base": "2020-10-09",
            "entra_no_total": True,
            "multa": {"aplicar": True, "percentual": "0.10", "base": "corrigido"},
            "honorarios": {"aplicar": True, "percentual": "0.10", "base": "corrigido"},
            "parcelas": [{"numero": 1}],
        },
        {
            "id": "parcelas", "nome": "parcelas fls. 3", "tipo": "valor",
            "valor": "14689.70", "data_base": "2022-02-06", "entra_no_total": True,
            "multa": {"aplicar": True, "percentual": "0.10", "base": "corrigido"},
            "honorarios": {"aplicar": True, "percentual": "0.10", "base": "corrigido"},
            "parcelas": [{"numero": 2}],
        },
    ]
    parcelas = [
        {
            "numero": 1, "verba_id": "entrada", "vencimento": "2020-10-09",
            "descricao": "Entrada", "principal": "36800.00", "total": "96383.41",
            "total_pago": "1881.86", "saldo": "94501.55",
            "correcao": {"aplicado": True, "correcao": "15430.17"},
            "juros": {"aplicado": True, "juros": "33707.20", "modalidade": "taxa_legal"},
            "multa": {"aplicado": True, "valor": "5223.02"},
            "honorarios": {"aplicado": True, "valor": "5223.02"},
            "pagamentos": [{"data": "2025-06-10", "valor": "1881.86", "descricao": "depósito parcial"}],
        },
        {
            "numero": 2, "verba_id": "parcelas", "vencimento": "2022-02-06",
            "descricao": "parcelas fls. 3", "principal": "14689.70", "total": "21905.27",
            "total_pago": "0.00", "saldo": "21905.27",
            "correcao": {"aplicado": True, "correcao": "3564.69"},
            "juros": {"aplicado": False, "juros": "0.00", "motivo": "Juros não aplicados (parâmetro explícito do usuário)."},
            "multa": {"aplicado": True, "valor": "1825.44"},
            "honorarios": {"aplicado": True, "valor": "1825.44"},
            "pagamentos": [],
        },
    ]
    return montar_memoria(
        codigo="VEC-2026-000001", tipo="parcelas", titulo="Vitor Santiago",
        processo="", parte="", app_name="Very Easy Calc", app_version="1.0.0",
        data_geracao=datetime(2026, 10, 15, tzinfo=timezone.utc), data_base="2026-09-30",
        configuracao={"correcao": {"aplicar": False}, "juros": {"aplicar": False}, "multa": {"aplicar": False}, "honorarios": {"aplicar": False}},
        payload={"pagamentos": []},
        resumo={"principal": "51489.70", "correcao": "18994.86", "juros": "33707.20", "multa": "7048.46", "honorarios": "7048.46", "total": "118288.68", "pagamentos": "1881.86", "saldo": "116406.82"},
        parcelas=parcelas, bases=[], avisos=[], verbas=verbas,
    )


def test_memoria_e_pdf_civel_refletem_encargos_por_verba_e_abatimento() -> None:
    memoria = _memoria_civel_com_encargos_e_abatimento()
    criterios = memoria["criterios"]

    assert criterios["multa"]["aplicar"] is True
    assert criterios["multa"]["por_verba"] is True
    assert criterios["honorarios"]["aplicar"] is True
    assert criterios["pagamentos_informados"] == memoria["parcelas"][0]["pagamentos"]
    linhas = dict(criterios_txt(criterios))
    assert linhas["Multa"].startswith("10% sobre corrigido")
    assert linhas["Honorários"].startswith("10% sobre corrigido")
    assert linhas["Pagamentos informados"] == "1"

    texto = _texto(montar_historia(memoria))
    assert "Fechamento" in texto
    assert "10/06/2025" in texto
    assert "depósito parcial" in texto
    assert "-R$ 1.881,86" in texto
    assert "Saldo final" in texto
    assert "R$ 116.406,82" in texto
    assert memoria["parcelas"][1]["juros"]["motivo"] == (
        "Juros não aplicados (parâmetro explícito do usuário)."
    )

    planilha = load_workbook(BytesIO(gerar_xlsx(memoria)), data_only=True)["Resumo"]
    valores = [linha for linha in planilha.iter_rows(values_only=True)]
    indice = next(i for i, linha in enumerate(valores) if linha[0] == "Abatimentos informados")
    assert valores[indice + 2][:4] == (
        "2025-06-10", -1881.86, None, "depósito parcial"
    )


def test_colunas_monetarias_do_resumo_civel_tem_largura_minima() -> None:
    memoria = _memoria_civel_com_encargos_e_abatimento()
    tabelas = []
    for item in montar_historia(memoria):
        if isinstance(item, Table):
            tabelas.append(item)
        elif isinstance(item, KeepTogether):
            tabelas.extend(x for x in item._content if isinstance(x, Table))

    # No A4 retrato, o resumo por verba foi substituído por tabelas de valores
    # principais e de composição/saldo para cada verba. As colunas monetárias
    # precisam preservar espaço suficiente para valores longos, sem ultrapassar
    # a área útil da página e provocar corte lateral.
    tabelas_monetarias = [
        tabela
        for tabela in tabelas
        if "— valores principais" in tabela._cellvalues[0][0].getPlainText()
        or "— composição e saldo" in tabela._cellvalues[0][0].getPlainText()
    ]
    assert len(tabelas_monetarias) == 4

    cabecalhos_monetarios = {"Principal", "Correção", "Juros", "Multa", "Hon.", "Pagto.", "Total", "Saldo"}
    for tabela in tabelas_monetarias:
        cabecalhos = [celula.getPlainText() for celula in tabela._cellvalues[1]]
        larguras = dict(zip(cabecalhos, tabela._colWidths))
        colunas_monetarias = [
            largura
            for cabecalho, largura in larguras.items()
            if cabecalho in cabecalhos_monetarios
        ]
        assert colunas_monetarias
        # R$ 999.999,99 cabe com margem interna; não há coluna monetária estreita.
        assert min(colunas_monetarias) >= 30 * mm
        assert sum(tabela._colWidths) <= LARGURA + 0.01


def test_exportacao_corrige_apresentacao_de_memoria_legada_sem_recalcular() -> None:
    memoria = deepcopy(_memoria_civel_com_encargos_e_abatimento())
    memoria["criterios"]["multa"] = {"aplicar": False}
    memoria["criterios"]["honorarios"] = {"aplicar": False}
    memoria["criterios"]["pagamentos_informados"] = []

    texto = _texto(montar_historia(memoria))
    assert "Multa\naplicada — configurado por verba; critérios distintos por verba" in texto
    assert "Honorários adicionais\naplicada — configurado por verba; critérios distintos por verba" in texto
    assert "Pagamentos informados\n1" in texto
    assert "depósito parcial" in texto

    planilha = load_workbook(BytesIO(gerar_xlsx(memoria)), data_only=True)["Resumo"]
    valores = [linha for linha in planilha.iter_rows(values_only=True)]
    assert any(linha[0] == "Abatimentos informados" for linha in valores)
